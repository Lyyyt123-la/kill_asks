"""把 docs/技术文档.md 渲染成排版好的 DOCX（封面 / 目录 / 标题 / 表格 / 图片 / 代码块 / 页眉页脚）。

用法：backend\\.venv\\Scripts\\python.exe scripts/build_tech_doc.py
输出：docs/技术文档_问数AskData.docx
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
SOURCE = DOCS / "技术文档.md"
OUTPUT = DOCS / "技术文档_问数AskData.docx"

GREEN = RGBColor(0x1F, 0x3D, 0x2E)
GREEN2 = RGBColor(0x2F, 0x6B, 0x4F)
GRAY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0x7A, 0x74, 0x68)

CN_BODY = "宋体"
CN_HEAD = "微软雅黑"
EN_FONT = "Times New Roman"
MONO = "Consolas"


def set_cn_font(run, font_name: str = CN_BODY) -> None:
    run.font.name = EN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), EN_FONT)
    rfonts.set(qn("w:hAnsi"), EN_FONT)


def shade(paragraph_or_cell, color: str) -> None:
    element = paragraph_or_cell._element
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    pr.append(shd)


def add_rich_text(paragraph, text: str, size: float | None = None, color: RGBColor | None = None) -> None:
    """支持 **加粗** 的简单行内格式。"""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        set_cn_font(run)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def style_setup(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x23, 0x21, 0x1D)
    normal.font.name = EN_FONT
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CN_BODY)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color, before, after in (
        ("Heading 1", 15, GREEN, 16, 6),
        ("Heading 2", 12, GREEN2, 10, 4),
        ("Heading 3", 11, RGBColor(0x3E, 0x3A, 0x32), 8, 3),
    ):
        style = document.styles[name]
        style.font.name = EN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CN_HEAD)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = EN_FONT
        style.font.size = Pt(10.5)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CN_BODY)
        style.paragraph_format.line_spacing = 1.4


def page_setup(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.2)

    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run("问数 AskData · 技术文档")
    set_cn_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = LIGHT

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_para.add_run("第 ")
    set_cn_font(run)
    run.font.size = Pt(8.5)
    # PAGE 域
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2 = footer_para.add_run()
    run2._element.append(fld_begin)
    run2._element.append(instr)
    run2._element.append(fld_end)
    run3 = footer_para.add_run(" 页")
    set_cn_font(run3)
    run3.font.size = Pt(8.5)


def cover(document: Document) -> None:
    for _ in range(3):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("问数 AskData")
    set_cn_font(run, CN_HEAD)
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = GREEN

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于自然语言交互的数据分析平台")
    set_cn_font(run, CN_HEAD)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x3E, 0x3A, 0x32)
    subtitle.paragraph_format.space_after = Pt(6)

    badge = document.add_paragraph()
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = badge.add_run("AI WEB 网页开发挑战赛 · 赛道二 · 技术文档")
    set_cn_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = GREEN2
    badge.paragraph_format.space_after = Pt(48)

    info = [
        ("学　　校", "广东科贸职业学院"),
        ("指导老师", "林龙键"),
        ("团队成员", "罗勇添、辛嘉炜、林俊宁"),
        ("提交日期", "2026 年 10 月"),
    ]
    table = document.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, (key, value) in enumerate(info):
        cells = table.rows[row_index].cells
        cells[0].width = Cm(3.6)
        cells[1].width = Cm(8.0)
        cell_para = cells[0].paragraphs[0]
        cell_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = cell_para.add_run(key + "：")
        set_cn_font(run)
        run.font.size = Pt(12)
        run.font.color.rgb = GRAY
        value_para = cells[1].paragraphs[0]
        run = value_para.add_run(value)
        set_cn_font(run)
        run.font.size = Pt(12)
        run.font.bold = True

    for _ in range(4):
        document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("本作品由项目组独立完成 · 截图与数据均来自系统真实运行结果")
    set_cn_font(run)
    run.font.size = Pt(9.5)
    run.font.color.rgb = LIGHT
    document.add_page_break()


def toc(document: Document, headings: list[str]) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("目　录")
    set_cn_font(run, CN_HEAD)
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = GREEN
    title.paragraph_format.space_after = Pt(14)

    for heading in headings:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(heading)
        set_cn_font(run)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x3E, 0x3A, 0x32)
    document.add_page_break()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.rows[row_index].cells[col_index]
            text = row[col_index] if col_index < len(row) else ""
            cell_para = cell.paragraphs[0]
            cell_para.paragraph_format.space_after = Pt(1)
            add_rich_text(cell_para, text, size=9)
            if row_index == 0:
                shade(cell, "E8F0EA")
                for run in cell_para.runs:
                    run.font.bold = True
                    run.font.color.rgb = GREEN
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(document: Document, relative: str, alt: str) -> None:
    image_path = (DOCS / relative).resolve()
    if not image_path.exists():
        document.add_paragraph(f"[图片缺失：{relative}]")
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(15.2))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(alt)
    set_cn_font(run)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = LIGHT
    caption.paragraph_format.space_after = Pt(8)


def add_code(document: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        shade(paragraph, "F5F2EA")
        run = paragraph.add_run(line if line else " ")
        run.font.name = MONO
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x2B, 0x2B, 0x27)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build() -> Path:
    raw = SOURCE.read_text(encoding="utf-8")
    lines = raw.splitlines()

    document = Document()
    style_setup(document)
    page_setup(document)
    cover(document)

    # 收集一级标题做目录
    headings = [line[3:].strip() for line in lines if line.startswith("## ") and len(line) > 3]
    toc(document, headings)

    # 跳过封面信息段（文件头到第一个 ---）
    start = 0
    for index, line in enumerate(lines):
        if line.strip() == "---":
            start = index + 1
            break

    index = start
    table_buffer: list[list[str]] = []
    code_buffer: list[str] = []
    in_code = False

    def flush_table() -> None:
        nonlocal table_buffer
        if table_buffer:
            add_table(document, table_buffer)
            table_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code(document, code_buffer)
                code_buffer = []
                in_code = False
            else:
                flush_table()
                in_code = True
            index += 1
            continue
        if in_code:
            code_buffer.append(line)
            index += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(set(cell) <= set("-: ") for cell in cells):
                index += 1
                continue
            table_buffer.append(cells)
            index += 1
            continue
        flush_table()

        if not stripped:
            index += 1
            continue
        if stripped == "---":
            document.add_paragraph().paragraph_format.space_after = Pt(0)
            index += 1
            continue

        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
        if image_match:
            add_image(document, image_match.group(2), image_match.group(1))
            index += 1
            continue

        if stripped.startswith("#### "):
            paragraph = document.add_paragraph(stripped[5:], style="Heading 3")
            set_cn_font(paragraph.runs[0], CN_HEAD)
            paragraph.runs[0].font.bold = True
            index += 1
            continue
        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_rich_text(paragraph, stripped[4:], size=12, color=GREEN2)
            index += 1
            continue
        if stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_rich_text(paragraph, stripped[3:], size=15, color=GREEN)
            index += 1
            continue
        if stripped.startswith("# "):
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            run = paragraph.add_run(stripped[2:])
            set_cn_font(run)
            run.font.size = Pt(10)
            run.font.italic = True
            run.font.color.rgb = GREEN2
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_rich_text(paragraph, stripped[2:])
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_rich_text(paragraph, stripped)
        index += 1

    flush_table()
    if in_code and code_buffer:
        add_code(document, code_buffer)

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print("docx ->", path, f"({path.stat().st_size / 1024:.0f} KB)")
